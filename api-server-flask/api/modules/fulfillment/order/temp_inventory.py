# -*- encoding: utf-8 -*-
"""
Scratch stock-on-hand for the DMS download step (`temp_inventory`).

The DMS operator uploads a stock sheet (PART# | QTY); building an order's DMS file then
allocates against it, so the file carries what can actually be supplied rather than what
was ordered. See schema.py for why the table is keyed on the part number string.

The upload is a full SNAPSHOT, not a delta: parts named in the sheet get that quantity,
and every part already on record but ABSENT from the sheet drops to 0. A part missing
from a stock report means "none in stock", which is not the same as "unchanged" — the
alternative would leave yesterday's quantity standing and over-promise on it.
"""

import csv
import io
from datetime import datetime, timedelta

import openpyxl

from api.shared.db_manager import mysql_manager
from api.core.logging import get_logger

logger = get_logger(__name__)


# How recent the last upload has to be for a DMS download to be allowed. Stock is only
# meaningful for as long as nobody has picked against it.
FRESHNESS_MINUTES = 30


class InventorySheetError(Exception):
    """The uploaded stock sheet could not be understood."""


def _norm(header):
    return ''.join(ch for ch in str(header or '').lower() if ch.isalnum())


# The quantity column, in priority order. EXACT matches are tried before anything fuzzy,
# and that ordering is load-bearing: the DMS stock export carries both 'Stock on Hand'
# and 'Actual Stock on Hand', so a substring test for "stock on hand" matches the wrong
# column just as readily as the right one.
_QTY_EXACT = (
    'stockonhand', 'stockinhand',
    'qty', 'quantity', 'stock', 'stockqty', 'availableqty', 'available',
    'onhand', 'closingstock', 'balance', 'freestock',
)


def _qty_column(header):
    """Index of the quantity column, or None."""
    for name in _QTY_EXACT:
        if name in header:
            return header.index(name)
    # Only now fall back to a contains test, for spellings like 'stockonhandnos'. Columns
    # qualified with another word ('actual…', 'virtual…') are skipped so the plain one is
    # still preferred when both exist.
    for idx, col in enumerate(header):
        if ('stockonhand' in col or 'stockinhand' in col) and not col.startswith(
                ('actual', 'virtual', 'foc', 'vor')):
            return idx
    return None


def _decode(raw):
    """Text from a delimited export, whatever it was saved as.

    The DMS stock export is UTF-16 with a BOM, which decoded as UTF-8 yields NUL bytes
    between every character and a header that matches nothing.
    """
    if raw[:2] in (b'\xff\xfe', b'\xfe\xff'):
        return raw.decode('utf-16')          # BOM picks the endianness
    for enc in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise InventorySheetError("the file's text encoding could not be read")


def _read_rows(file_storage):
    """Rows of cells from either an Excel workbook or a delimited text export.

    Dispatches on the file's CONTENT, not its name: the stock export arrives named .csv
    but is tab-separated, and operators rename these files freely.
    """
    raw = file_storage.read()
    if not raw:
        raise InventorySheetError("the file is empty")

    if raw[:2] == b'PK':                      # xlsx/xlsm are zip archives
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        except Exception as e:
            raise InventorySheetError(f"could not read the file as an Excel workbook ({e})")
        return [list(r) for r in wb.active.iter_rows(values_only=True)]

    text = _decode(raw)
    first = text.split('\n', 1)[0]
    # Sniff the separator off the header rather than assuming a comma — this export is
    # tab-separated despite the .csv name.
    delimiter = max(('\t', ',', ';', '|'), key=first.count)
    if first.count(delimiter) == 0:
        raise InventorySheetError(
            "the file has only one column — expected a part number and a quantity")
    # csv handles the quoting ("K06431KTNA701S") that a plain split would leave behind.
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter) if row]


def parse_inventory_sheet(file_storage):
    """Parse a stock export into [{part_number, quantity}].

    The part number is taken from the FIRST column, whatever it is called — the export's
    own header is 'Part #' but that is not worth depending on, and every stock report seen
    so far leads with the part. The quantity comes from the named 'Stock on Hand' column
    (see _qty_column), since its position varies between reports.

    Unlike parse_part_convertor, a quantity of 0 is KEPT: on a stock sheet zero is a real
    statement about a part, not a row to skip.
    """
    rows = _read_rows(file_storage)
    if not rows:
        raise InventorySheetError("the sheet is empty")

    header = [_norm(c) for c in rows[0]]
    i_qty = _qty_column(header)
    if i_qty is None:
        raise InventorySheetError(
            "could not find a stock quantity column — expected one named "
            "'Stock on Hand' (or Qty / Quantity). Found: "
            + ', '.join(str(c) for c in rows[0][:12] if str(c or '').strip()))
    if len(header) < 2:
        raise InventorySheetError("the sheet needs at least two columns")

    def cell(row, idx):
        return row[idx] if idx < len(row) else None

    # Last value wins if a part is listed twice — the table's unique key cannot hold both,
    # and silently summing them would invent stock the sheet never claimed.
    seen = {}
    for row in rows[1:]:
        if not row:
            continue
        part = cell(row, 0)                   # first column, as specified
        part = str(part).strip().strip('"') if part is not None else ''
        if not part:
            continue
        raw = cell(row, i_qty)
        if raw is None or str(raw).strip() == '':
            qty = 0
        else:
            try:
                # Thousands separators and stray quotes both appear in these exports.
                qty = int(float(str(raw).strip().strip('"').replace(',', '')))
            except (TypeError, ValueError):
                qty = 0
        # Negative stock is meaningless downstream (it would read as supply); floor it.
        seen[part] = max(0, qty)

    if not seen:
        raise InventorySheetError(
            "no usable rows found (each needs a part number and a quantity)")
    return [{'part_number': p, 'quantity': q} for p, q in seen.items()]


# Rows per statement for the bulk writes. Big enough that a 20k-part sheet is ~20 round
# trips rather than 20k, small enough that no single packet gets unwieldy.
_CHUNK = 1000


def replace_stock(items):
    """Apply an uploaded snapshot. Returns (updated, inserted, zeroed).

    Written in bulk, deliberately: these sheets run to tens of thousands of parts, and a
    statement per part (with a per-row subselect to resolve the catalogue id) is the N+1
    the coding standards call out — it made a 20k-row upload take minutes. The work is now
    a handful of statements whose cost barely moves with sheet size.

    One transaction, so a download can never allocate against a half-applied sheet.
    """
    now = datetime.utcnow()
    uploaded = {it['part_number']: it['quantity'] for it in items}
    parts = list(uploaded)

    # The three counts the operator is shown, derived from one read of the key column
    # instead of from per-row rowcounts.
    existing = {r['part_number'] for r in (mysql_manager.execute_query(
        "SELECT part_number FROM temp_inventory") or [])}
    inserted = sum(1 for p in parts if p not in existing)
    updated = len(parts) - inserted
    zeroed = sum(1 for p in existing if p not in uploaded)

    # Catalogue ids for the whole sheet in a few queries rather than one per part.
    # product_string is UNIQUE, so this is an index lookup per chunk.
    pid = {}
    for i in range(0, len(parts), _CHUNK):
        chunk = parts[i:i + _CHUNK]
        ph = ",".join(["%s"] * len(chunk))
        for r in (mysql_manager.execute_query(
                f"SELECT product_id, product_string FROM product "
                f"WHERE product_string IN ({ph})", tuple(chunk)) or []):
            pid[r['product_string']] = r['product_id']

    with mysql_manager.get_cursor() as cursor:
        # A snapshot means every part not named in the sheet is zero. Zeroing the whole
        # table in one statement and then writing the sheet's numbers over it gets there
        # in two steps; the alternative, NOT IN (<20k placeholders>), is a statement
        # megabytes long that the server has to parse before it can do anything.
        cursor.execute(
            "UPDATE temp_inventory SET quantity = 0, uploaded_at = %s", (now,))
        for i in range(0, len(parts), _CHUNK):
            chunk = parts[i:i + _CHUNK]
            args = []
            for p in chunk:
                args += [p, pid.get(p), uploaded[p], now, now]
            placeholders = ",".join(["(%s,%s,%s,%s,%s)"] * len(chunk))
            # uq_temp_inv_part is what makes this an upsert: a part already on record has
            # its quantity replaced, a new one is inserted.
            cursor.execute(
                "INSERT INTO temp_inventory "
                "  (part_number, product_id, quantity, uploaded_at, created_at) "
                f"VALUES {placeholders} "
                "ON DUPLICATE KEY UPDATE quantity = VALUES(quantity), "
                "  product_id = VALUES(product_id), uploaded_at = VALUES(uploaded_at)",
                tuple(args))

    logger.info("temp_inventory snapshot applied", extra={
        'updated': updated, 'inserted': inserted, 'zeroed': zeroed,
        'parts_in_sheet': len(parts)})
    return updated, inserted, zeroed


def last_uploaded_at():
    """When stock was last uploaded, or None if it never has been."""
    row = mysql_manager.execute_query(
        "SELECT MAX(uploaded_at) AS t FROM temp_inventory")
    return row[0]['t'] if row else None


def status():
    """Freshness summary for the UI: enough to render the banner and pre-warn before a
    download is attempted."""
    last = last_uploaded_at()
    counts = mysql_manager.execute_query(
        "SELECT COUNT(*) AS parts, COALESCE(SUM(quantity), 0) AS units FROM temp_inventory")
    parts = counts[0]['parts'] if counts else 0
    units = int(counts[0]['units'] or 0) if counts else 0
    age_min = None
    if last is not None:
        age_min = max(0, int((datetime.utcnow() - last).total_seconds() // 60))
    return {
        'last_uploaded_at': last.isoformat() if last is not None else None,
        'age_minutes': age_min,
        'fresh': (age_min is not None and age_min < FRESHNESS_MINUTES),
        'freshness_minutes': FRESHNESS_MINUTES,
        'parts': parts,
        'units': units,
    }


def staleness_error():
    """The reason a download must be refused, or None when stock is fresh enough.

    Returned as a message rather than raised so the caller keeps control of the status
    code, and so the wording reaches the operator unchanged.
    """
    last = last_uploaded_at()
    if last is None:
        return ("No inventory has been uploaded yet. Upload the current stock sheet "
                "before downloading a DMS file.")
    age = datetime.utcnow() - last
    if age > timedelta(minutes=FRESHNESS_MINUTES):
        mins = int(age.total_seconds() // 60)
        return (f"Inventory was last uploaded {mins} minutes ago, which is older than "
                f"{FRESHNESS_MINUTES} minutes. Upload the current stock sheet before "
                f"downloading a DMS file.")
    return None


def allocate(order_items):
    """Decide what can be supplied for `order_items`, and take it out of stock.

    Each item needs sku_code and quantity. Returns (allocated, shortfalls) where
    `allocated` mirrors the input with `quantity` cut to what stock could cover, and
    `shortfalls` lists {part_number, ordered, supplied} for anything cut — so the caller
    can tell the operator what was trimmed rather than handing over a quietly smaller file.

    Lines that can be supplied at all are kept; a line with nothing in stock is dropped,
    because a DMS row asking for zero pieces is not an order for anything.

    Reads and writes inside one transaction: two operators downloading at the same moment
    must not both be promised the last piece.
    """
    codes = [it['sku_code'] for it in order_items if it.get('sku_code')]
    if not codes:
        return [], []

    allocated, shortfalls = [], []
    with mysql_manager.get_cursor() as cursor:
        ph = ",".join(["%s"] * len(codes))
        # FOR UPDATE so a concurrent download waits rather than reading the same stock.
        cursor.execute(
            f"""SELECT part_number, quantity FROM temp_inventory
                 WHERE part_number IN ({ph}) FOR UPDATE""", tuple(codes))
        on_hand = {r['part_number']: int(r['quantity'] or 0) for r in cursor.fetchall()}

        taken = {}
        for it in order_items:
            part = it.get('sku_code')
            ordered = int(it.get('quantity') or 0)
            if not part or ordered <= 0:
                continue
            # A part with no row at all is out of stock, same as a row reading 0.
            available = on_hand.get(part, 0) - taken.get(part, 0)
            supplied = max(0, min(ordered, available))
            if supplied < ordered:
                shortfalls.append({'part_number': part, 'ordered': ordered,
                                   'supplied': supplied})
            if supplied <= 0:
                continue
            taken[part] = taken.get(part, 0) + supplied
            line = dict(it)
            line['quantity'] = supplied
            allocated.append(line)

        for part, qty in taken.items():
            # Guarded with GREATEST so a race can never drive stock negative.
            cursor.execute(
                """UPDATE temp_inventory
                      SET quantity = GREATEST(0, quantity - %s)
                    WHERE part_number = %s""", (qty, part))

    return allocated, shortfalls
